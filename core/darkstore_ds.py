"""Report Daily Performance DS & MTD DS Performance + Delivery.

Sumber: notebooks/daily/report_daily_performance_ds.ipynb dan
notebooks/mtd/mtd_ds_performance_delivery.ipynb. Kedua notebook berbagi
logika yang sama; versi MTD menambah file Report Summary (delivery) dan
kolom %Ontime / %Late di sheet Summary.

PERUBAHAN (fix KeyError JHK/SALES/... saat ZIP berisi 2 Detail Data + 2 OOS):
- find_input_files: pilih Detail Data yang TRX (bukan SLA) dan OOS periode (Full/MTD).
- load_detail_data: validasi kolom wajib, pesan error jelas kalau salah file.
"""
import os

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter

from .common import Log, recalc_with_libreoffice

BULAN_ABBR_ID = {1: 'Jan', 2: 'Feb', 3: 'Mar', 4: 'Apr', 5: 'Mei', 6: 'Jun',
                 7: 'Jul', 8: 'Agu', 9: 'Sept', 10: 'Okt', 11: 'Nov', 12: 'Des'}
BULAN_FULL_ID = {1: 'Januari', 2: 'Februari', 3: 'Maret', 4: 'April', 5: 'Mei', 6: 'Juni',
                 7: 'Juli', 8: 'Agustus', 9: 'September', 10: 'Oktober', 11: 'November', 12: 'Desember'}

N_FALLBACK_DAYS = 2

# Kolom wajib pada file Detail Data TRX (file SLA tidak punya kolom-kolom ini)
REQUIRED_DETAIL_COLS = ["TANGGAL", "KD_STORE", "NAMA_STORE", "NAMA_BRANCH",
                        "JHK", "SALES", "SALES_TAGI", "SPD", "STD", "APC",
                        "PERCENT_GM", "PCT_OOS_OFMB"]

FILL_HEADER = PatternFill("solid", fgColor="D9E1F2")
FONT_TITLE = Font(bold=True, size=14)
FONT_SUB = Font(size=11)
FONT_HEADER = Font(bold=True, size=12)
FONT_HEADER_SM = Font(bold=True, size=11)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="000000")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FMT_INT = "#,##0"
FMT_PCT = "0.00%"


# ---------------------------------------------------------------- input
def find_input_files(paths, need_report_summary=False):
    """Cari Detail Data TRX (.csv), [OOS] By Toko (.xlsx) dan (opsional) Report Summary (.xlsx).

    Aturan pemilihan (berdasarkan nama file):
    - CSV  : utamakan yang mengandung 'trx'; file 'sla' tidak dipakai di sini.
    - OOS  : utamakan yang mengandung 'full' atau 'mtd' (OOS periode 1-tanggal akhir),
             kalau tidak ada pakai OOS pertama yang ditemukan.
    """
    names = {p: os.path.basename(p).lower() for p in paths}
    csvs = [p for p in paths if names[p].endswith(".csv")]

    detail = [p for p in csvs if "detail" in names[p] and "data" in names[p]] or csvs
    csv = (next((p for p in detail if "trx" in names[p]), None)
           or next((p for p in detail if "sla" not in names[p]), None))

    xlsxs = [p for p in paths if names[p].endswith(".xlsx") and "report daily" not in names[p]]
    oos_all = [p for p in xlsxs if "oos" in names[p]]
    oos = (next((p for p in oos_all if "full" in names[p] or "mtd" in names[p]), None)
           or (oos_all[0] if oos_all else None))

    report_summary = None
    if need_report_summary:
        rest = [p for p in xlsxs if p != oos]
        report_summary = next((p for p in rest if "summary" in names[p] or "report" in names[p]),
                              rest[0] if rest else None)
    if oos is None:
        rest = [p for p in xlsxs if p != report_summary]
        oos = rest[0] if rest else None

    if csv is None:
        raise FileNotFoundError(
            "File Detail Data TRX (.csv) tidak ditemukan. "
            "Pastikan ada file CSV yang namanya mengandung 'TRX' (bukan 'SLA').")
    if oos is None:
        raise FileNotFoundError("File [OOS] By Toko (.xlsx) tidak ditemukan.")
    if need_report_summary and report_summary is None:
        raise FileNotFoundError("File Report Summary Dashboard (.xlsx) tidak ditemukan.")
    return csv, oos, report_summary


def load_detail_data(csv_path, kd_store_as_str=False):
    df = pd.read_csv(csv_path)
    df.columns = [str(c).replace("\ufeff", "").strip() for c in df.columns]
    missing = [c for c in REQUIRED_DETAIL_COLS if c not in df.columns]
    if missing:
        raise ValueError(
            f"File '{os.path.basename(csv_path)}' bukan Detail Data TRX. "
            f"Kolom hilang: {missing}. Kolom yang terbaca: {list(df.columns)[:15]}...")
    df["TANGGAL"] = pd.to_datetime(df["TANGGAL"])
    if kd_store_as_str:
        df["KD_STORE"] = df["KD_STORE"].astype(str).str.strip()
    return df


def load_oos_data(xlsx_path):
    df = pd.read_excel(xlsx_path, sheet_name=0)
    df.columns = [str(c).strip() for c in df.columns]
    df = df[df.iloc[:, 0].astype(str).str.strip().str.lower() != "grand total"]
    df = df.dropna(subset=[df.columns[0]]).reset_index(drop=True)
    needed = ["% OOS OFMB", "% OOS TAG I", "% OOS TAG K"]
    missing = [c for c in needed if c not in df.columns]
    if missing:
        raise ValueError(
            f"File OOS '{os.path.basename(xlsx_path)}' kolom hilang: {missing}. "
            f"Kolom yang terbaca: {list(df.columns)}")
    return df


def _read_report_summary_raw(xlsx_path):
    df = pd.read_excel(xlsx_path, sheet_name=0)
    df.columns = [str(c).strip() for c in df.columns]
    required = ["KD_STORE", "JUMLAH_DELIVERY", "DELIVERY_ONTIME", "DELIVERY_LATE"]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise KeyError(f"Kolom Report Summary tidak ditemukan: {missing}. "
                       f"Kolom yang tersedia: {list(df.columns)}")
    df["KD_STORE"] = df["KD_STORE"].astype(str).str.strip()
    for col in ["JUMLAH_DELIVERY", "DELIVERY_ONTIME", "DELIVERY_LATE"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df["%Ontime"] = np.where(df["JUMLAH_DELIVERY"] > 0, df["DELIVERY_ONTIME"] / df["JUMLAH_DELIVERY"], 0)
    df["%Late"] = np.where(df["JUMLAH_DELIVERY"] > 0, df["DELIVERY_LATE"] / df["JUMLAH_DELIVERY"], 0)
    return df


def load_report_summary(xlsx_path):
    df = _read_report_summary_raw(xlsx_path)
    return (df[["KD_STORE", "%Ontime", "%Late"]]
            .drop_duplicates(subset="KD_STORE", keep="last")
            .reset_index(drop=True))


# ---------------------------------------------------------------- agregasi
def sort_by_net_sales(g, sales_col="Sum of SALES"):
    if sales_col not in g.columns:
        return g.reset_index(drop=True)
    return g.sort_values(sales_col, ascending=False, na_position="last").reset_index(drop=True)


def daily_summary(df, tanggal):
    sub = df[df["TANGGAL"] == tanggal]
    g = sub.groupby(["KD_STORE", "NAMA_STORE", "NAMA_BRANCH"], as_index=False).agg(**{
        "Sum of JHK": ("JHK", "sum"),
        "Average of SPD": ("SPD", "mean"),
        "SPD Tag I": ("SALES_TAGI", "sum"),
        "Average of STD": ("STD", "mean"),
        "Average of APC": ("APC", "mean"),
        "Average of PERCENT_GM": ("PERCENT_GM", "mean"),
        "__sales__": ("SALES", "sum"),
    })
    return (g.sort_values("__sales__", ascending=False, na_position="last")
             .drop(columns="__sales__").reset_index(drop=True))


_FULL_AGG = {
    "Sum of JHK": ("JHK", "sum"),
    "Sum of SALES": ("SALES", "sum"),
    "Sum of SALES_TAGI": ("SALES_TAGI", "sum"),
    "Average of SPD": ("SPD", "mean"),
    "Average of STD": ("STD", "mean"),
    "Average of APC": ("APC", "mean"),
    "Average of PERCENT_GM": ("PERCENT_GM", "mean"),
    "Average of PCT_OOS_OFMB": ("PCT_OOS_OFMB", "mean"),
}


def daily_summary_full(df, tanggal):
    sub = df[df["TANGGAL"] == tanggal]
    g = sub.groupby(["KD_STORE", "NAMA_STORE", "NAMA_BRANCH"], as_index=False).agg(**_FULL_AGG)
    return sort_by_net_sales(g)


def period_summary(df):
    g = df.groupby(["KD_STORE", "NAMA_STORE", "NAMA_BRANCH"], as_index=False).agg(**_FULL_AGG)
    return sort_by_net_sales(g)


def sheet_label(tanggal):
    ts = pd.Timestamp(tanggal)
    return f"Data {ts.day} {BULAN_ABBR_ID[ts.month]} {ts.year}"


def periode_title(start, end):
    start, end = pd.Timestamp(start), pd.Timestamp(end)
    if (start.year, start.month) == (end.year, end.month):
        return f"Periode {start.day} - {end.day} {BULAN_FULL_ID[start.month]} {start.year}"
    return (f"Periode {start.day} {BULAN_FULL_ID[start.month]} {start.year} - "
            f"{end.day} {BULAN_FULL_ID[end.month]} {end.year}")


def build_summary_table(df_detail, df_oos, df_report_summary=None, n_fallback_days=N_FALLBACK_DAYS):
    all_dates = sorted(df_detail["TANGGAL"].unique(), reverse=True)
    fallback_dates = all_dates[:n_fallback_days]
    daily_tables = {pd.Timestamp(d): daily_summary(df_detail, d) for d in fallback_dates}
    df_period = period_summary(df_detail)

    master = (df_detail.sort_values("TANGGAL")
              .drop_duplicates(subset="KD_STORE", keep="last")
              [["KD_STORE", "NAMA_STORE", "NAMA_BRANCH"]])
    summary = master.merge(df_period, on=["KD_STORE", "NAMA_STORE", "NAMA_BRANCH"], how="left")

    source_sheet, missing_stores = [], []
    for kd in summary["KD_STORE"]:
        for d in fallback_dates:
            if (daily_tables[pd.Timestamp(d)]["KD_STORE"] == kd).any():
                source_sheet.append(sheet_label(d))
                break
        else:
            source_sheet.append(None)
            missing_stores.append(kd)
    summary["__source_sheet__"] = source_sheet

    oos = df_oos.rename(columns={df_oos.columns[0]: "Kode Toko"})
    oos_small = oos[["Kode Toko", "% OOS OFMB", "% OOS TAG I", "% OOS TAG K"]].copy()
    summary = summary.merge(oos_small, left_on="KD_STORE", right_on="Kode Toko", how="left")

    if df_report_summary is not None:
        summary = summary.merge(df_report_summary[["KD_STORE", "%Ontime", "%Late"]],
                                on="KD_STORE", how="left")

    summary = sort_by_net_sales(summary)
    summary.insert(0, "No", range(1, len(summary) + 1))
    return summary, daily_tables, df_period, missing_stores, fallback_dates


# ---------------------------------------------------------------- penulisan Excel
def style_header_row(ws, row, col_start, col_end, font=FONT_HEADER_SM):
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = font
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_ALL


def autosize(ws, ncols, min_width=10, max_width=32):
    for c in range(1, ncols + 1):
        letter = get_column_letter(c)
        best = min_width
        for cell in ws[letter]:
            if cell.value is not None:
                best = max(best, min(max_width, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = best


def write_data_sheet(wb, df_detail):
    ws = wb.create_sheet("Data")
    ws.append(list(df_detail.columns))
    style_header_row(ws, 1, 1, len(df_detail.columns))
    for row in df_detail.itertuples(index=False):
        ws.append(list(row))
    date_col = list(df_detail.columns).index("TANGGAL") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=date_col).number_format = "m/d/yy h:mm"
    ws.freeze_panes = "A2"
    autosize(ws, len(df_detail.columns))


def write_daily_tables(wb, daily_tables):
    for tanggal, tbl in daily_tables.items():
        ws = wb.create_sheet(sheet_label(tanggal))
        ws.append(list(tbl.columns))
        style_header_row(ws, 1, 1, len(tbl.columns))
        for row in tbl.itertuples(index=False):
            ws.append(list(row))
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=5).number_format = FMT_INT
            ws.cell(row=r, column=6).number_format = FMT_INT
            ws.cell(row=r, column=8).number_format = FMT_INT
            ws.cell(row=r, column=9).number_format = FMT_PCT
        autosize(ws, len(tbl.columns))


def write_pivot_sheet(wb, df_period, df_detail, latest_date):
    ws = wb.create_sheet("Pivot Table")
    left = daily_summary_full(df_detail, latest_date)
    right = df_period

    ws.cell(row=1, column=1,
            value=f"Filter Tanggal: {pd.Timestamp(latest_date).strftime('%d/%m/%Y')}").font = Font(bold=True)
    headers = list(left.columns)
    ncol_left = len(headers)
    for j, h in enumerate(headers, start=1):
        ws.cell(row=2, column=j, value=h)
    style_header_row(ws, 2, 1, ncol_left)
    for i, row in enumerate(left.itertuples(index=False), start=3):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)

    start_col_right = ncol_left + 3
    ws.cell(row=1, column=start_col_right, value="Filter Tanggal: (All)").font = Font(bold=True)
    for j, h in enumerate(headers, start=start_col_right):
        ws.cell(row=2, column=j, value=h)
    style_header_row(ws, 2, start_col_right, start_col_right + ncol_left - 1)
    for i, row in enumerate(right.itertuples(index=False), start=3):
        for j, v in enumerate(row, start=start_col_right):
            ws.cell(row=i, column=j, value=v)

    for block_start in (1, start_col_right):
        for r in range(3, 3 + max(len(left), len(right))):
            for j, h in enumerate(headers, start=block_start):
                if h.startswith(("Sum of SALES", "Average of SPD", "Average of STD", "Average of APC")):
                    ws.cell(row=r, column=j).number_format = FMT_INT
                elif h.startswith(("Average of PERCENT_GM", "Average of PCT_OOS")):
                    ws.cell(row=r, column=j).number_format = FMT_PCT
    autosize(ws, start_col_right + ncol_left - 1)


def write_oos_sheet(wb, df_oos):
    ws = wb.create_sheet("OOS")
    ws.append(list(df_oos.columns))
    style_header_row(ws, 1, 1, len(df_oos.columns))
    for row in df_oos.itertuples(index=False):
        ws.append(list(row))
    for r in range(2, ws.max_row + 1):
        for c in range(4, len(df_oos.columns) + 1):
            ws.cell(row=r, column=c).number_format = FMT_PCT
    autosize(ws, len(df_oos.columns))


def write_report_summary_sheet(wb, report_summary_path):
    df = _read_report_summary_raw(report_summary_path)
    ws = wb.create_sheet("Report Summary")
    headers = ["KD_STORE", "JUMLAH_DELIVERY", "DELIVERY_ONTIME", "DELIVERY_LATE", "%Ontime", "%Late"]
    ws.append(headers)
    style_header_row(ws, 1, 1, 6)
    for row in df[headers].itertuples(index=False, name=None):
        ws.append(list(row))
    for r in range(2, ws.max_row + 1):
        for c in (2, 3, 4):
            ws.cell(r, c).number_format = FMT_INT
        for c in (5, 6):
            ws.cell(r, c).number_format = FMT_PCT
        for c in range(1, 7):
            ws.cell(r, c).border = BORDER_ALL
            ws.cell(r, c).alignment = ALIGN_CENTER
    ws.freeze_panes = "A2"
    autosize(ws, 6)


def write_summary_sheet(wb, summary, start_date, end_date, with_delivery=False):
    ws = wb.create_sheet("Summary")
    last_col = 24 if with_delivery else 22

    ws.cell(row=3, column=2, value="Daily Report Performance Darkstore").font = FONT_TITLE
    ws.cell(row=4, column=2, value=periode_title(start_date, end_date)).font = FONT_SUB

    end_ts = pd.Timestamp(end_date)
    label_daily = f"Sales Darkstore {end_ts.day} {BULAN_FULL_ID[end_ts.month]} {end_ts.year}"
    label_period = f"Sales Darkstore {periode_title(start_date, end_date).replace('Periode ', '')}"

    h1, h2 = 5, 6
    for col, val in ((2, "No"), (3, "Kode Toko"), (4, "Nama Toko"), (5, "Cabang"),
                     (6, label_daily), (12, label_period)):
        ws.cell(row=h1, column=col, value=val)
    for c in range(2, 6):
        ws.merge_cells(start_row=h1, start_column=c, end_row=h2, end_column=c)
    ws.merge_cells(start_row=h1, start_column=6, end_row=h1, end_column=11)
    ws.merge_cells(start_row=h1, start_column=12, end_row=h1, end_column=last_col)

    daily_jhk = "JHK" if with_delivery else " JHK"  # sama seperti notebook masing-masing
    sub_headers = [daily_jhk, "SPD", "SPD Tag I", "STD", "APC", "%GM",
                   "JHK", "Total Net Sales", "Total Net Sales Tag I", "SPD", "SPD Tag I",
                   "STD", "APC", "%GM", "%OOS OFMB", "%OOS I", "%OOS K"]
    if with_delivery:
        sub_headers += ["%Ontime", "%Late"]
    for j, h in enumerate(sub_headers, start=6):
        ws.cell(row=h2, column=j, value=h)
    style_header_row(ws, h1, 2, last_col, font=FONT_HEADER)
    style_header_row(ws, h2, 2, last_col, font=FONT_HEADER)

    int_cols = [6, 7, 8, 9, 10, 12, 13, 14, 15, 16, 17, 18] if with_delivery \
        else [6, 7, 8, 9, 10, 13, 14, 15, 16, 17, 18]
    pct_cols = [11, 19, 20, 21, 22] + ([23, 24] if with_delivery else [])

    first = 7
    records = summary.to_dict("records")
    for r, d in enumerate(records, start=first):
        ws.cell(row=r, column=2, value=d["No"])
        ws.cell(row=r, column=3, value=d["KD_STORE"])
        ws.cell(row=r, column=4, value=d["NAMA_STORE"])
        ws.cell(row=r, column=5, value=d["NAMA_BRANCH"])

        src = d["__source_sheet__"]
        if src is not None:
            sep = "," if with_delivery else ", "
            for k, col in enumerate(range(6, 12), start=4):
                ws.cell(row=r, column=col, value=f"=VLOOKUP(C{r},'{src}'!A:I{sep}{k},FALSE)")

        ws.cell(row=r, column=12, value=d["Sum of JHK"])
        ws.cell(row=r, column=13, value=d["Sum of SALES"])
        ws.cell(row=r, column=14, value=d["Sum of SALES_TAGI"])
        ws.cell(row=r, column=15, value=d["Average of SPD"])
        ws.cell(row=r, column=16, value=f"=N{r}/L{r}")
        ws.cell(row=r, column=17, value=d["Average of STD"])
        ws.cell(row=r, column=18, value=d["Average of APC"])
        ws.cell(row=r, column=19, value=d["Average of PERCENT_GM"])
        ws.cell(row=r, column=20, value=d.get("% OOS OFMB"))
        ws.cell(row=r, column=21, value=d.get("% OOS TAG I"))
        ws.cell(row=r, column=22, value=d.get("% OOS TAG K"))
        if with_delivery:
            ws.cell(row=r, column=23,
                    value=f"=IFERROR(VLOOKUP(C{r},'Report Summary'!$A:$F,5,FALSE),0)")
            ws.cell(row=r, column=24,
                    value=f"=IFERROR(VLOOKUP(C{r},'Report Summary'!$A:$F,6,FALSE),0)")

        for c in range(2, last_col + 1):
            ws.cell(row=r, column=c).border = BORDER_ALL
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center")
        for c in int_cols:
            ws.cell(row=r, column=c).number_format = FMT_INT
        for c in pct_cols:
            ws.cell(row=r, column=c).number_format = FMT_PCT
        if not with_delivery:
            ws.cell(row=r, column=12).number_format = "0"

    total_row = first + len(records)
    ws.cell(row=total_row, column=2, value="Average/Sum ")
    ws.merge_cells(start_row=total_row, start_column=2, end_row=total_row, end_column=5)
    avg_cols = ["F", "G", "H", "I", "J", "K", "L", "O", "P", "Q", "R", "S", "T", "U", "V"]
    if with_delivery:
        avg_cols += ["W", "X"]
    for col_letter in avg_cols:
        ws.cell(row=total_row, column=column_index_from_string(col_letter),
                value=f"=AVERAGE({col_letter}{first}:{col_letter}{total_row - 1})")
    for col_letter in ("M", "N"):
        ws.cell(row=total_row, column=column_index_from_string(col_letter),
                value=f"=SUM({col_letter}{first}:{col_letter}{total_row - 1})")
    for c in range(2, last_col + 1):
        ws.cell(row=total_row, column=c).font = Font(bold=True)
        ws.cell(row=total_row, column=c).border = BORDER_ALL
    if with_delivery:
        for c in int_cols:
            ws.cell(row=total_row, column=c).number_format = FMT_INT
        for c in pct_cols:
            ws.cell(row=total_row, column=c).number_format = FMT_PCT

    ws.freeze_panes = "F7" if with_delivery else ws.cell(row=first, column=1).coordinate
    autosize(ws, last_col)
    ws.column_dimensions["B"].width = 6


# ---------------------------------------------------------------- entry point
def run(paths, output_dir, with_delivery=False, report_summary_path=None,
        n_fallback_days=N_FALLBACK_DAYS, recalc=True, log=None):
    """Proses file input -> path file Excel hasil.

    paths: daftar file (sudah diekstrak dari ZIP).
    with_delivery: True = versi MTD (butuh file Report Summary Dashboard).
    """
    log = log or Log()
    if report_summary_path:
        paths = [p for p in paths if p != report_summary_path]
    csv_path, oos_path, rs_path = find_input_files(
        paths, need_report_summary=with_delivery and not report_summary_path)
    rs_path = report_summary_path or rs_path
    log("Detail Data   :", os.path.basename(csv_path))
    log("OOS By Toko   :", os.path.basename(oos_path))
    if with_delivery:
        log("Report Summary:", os.path.basename(rs_path))

    df_detail = load_detail_data(csv_path, kd_store_as_str=with_delivery)
    df_oos = load_oos_data(oos_path)
    df_rs = load_report_summary(rs_path) if with_delivery else None
    log(f"Detail Data: {len(df_detail)} baris, {df_detail['KD_STORE'].nunique()} toko, periode "
        f"{df_detail['TANGGAL'].min().date()} s.d. {df_detail['TANGGAL'].max().date()}")
    log(f"OOS By Toko: {len(df_oos)} toko")
    from .storage import save_if_configured
    save_if_configured(df_detail, df_oos, log)
    try:
        from .storage import save_ds_sla_daily_if_configured
        _sla_csv = next((p for p in paths if str(p).lower().endswith('.csv')
                         and 'sla' in os.path.basename(p).lower()), None)
        if _sla_csv:
            log('File SLA DS  :', os.path.basename(_sla_csv))
            save_ds_sla_daily_if_configured(_sla_csv, log)
        else:
            log('[Penyimpanan] OTD/LATE harian DS dilewati (tidak ada file SLA .csv di upload)')
    except Exception as e:  # noqa: BLE001
        log(f'[Penyimpanan] DS SLA harian GAGAL, report tetap dibuat: {type(e).__name__}: {e}')
    if with_delivery:
        try:
            from .storage import save_ds_sla_if_configured
            save_ds_sla_if_configured(_read_report_summary_raw(rs_path), df_detail["TANGGAL"].max(), log)
        except Exception as e:  # noqa: BLE001
            log(f"[Penyimpanan] DS SLA GAGAL, report tetap dibuat: {type(e).__name__}: {e}")

    summary, daily_tables, df_period, missing, fallback_dates = build_summary_table(
        df_detail, df_oos, df_rs, n_fallback_days)
    log(f"Total toko di Summary: {len(summary)}")
    if missing:
        log(f"[Peringatan] {len(missing)} toko tidak ada data di {n_fallback_days} tanggal "
            f"terakhir, kolom Daily dikosongkan: {missing}")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_data_sheet(wb, df_detail)
    write_daily_tables(wb, daily_tables)
    write_pivot_sheet(wb, df_period, df_detail, latest_date=fallback_dates[0])
    write_oos_sheet(wb, df_oos)
    if with_delivery:
        write_report_summary_sheet(wb, rs_path)
    write_summary_sheet(wb, summary, df_detail["TANGGAL"].min(), df_detail["TANGGAL"].max(),
                        with_delivery=with_delivery)

    order = (["Data"] + [sheet_label(d) for d in sorted(fallback_dates)] + ["Pivot Table", "OOS"]
             + (["Report Summary"] if with_delivery else []) + ["Summary"])
    wb._sheets = [wb[name] for name in order]

    name = "Report MTD DS Performance Delivery.xlsx" if with_delivery else "Report Daily DS.xlsx"
    out = os.path.join(output_dir, name)
    wb.save(out)
    log("Tersimpan:", name, "| Sheet:", ", ".join(wb.sheetnames))
    if recalc:
        recalc_with_libreoffice(out, log=log)
    preview = summary.drop(columns=["__source_sheet__", "Kode Toko"], errors="ignore")
    return out, preview, []
