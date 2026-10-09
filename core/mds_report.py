"""Report Daily MDS (Mini Darkstore).

Input (boleh 1 ZIP berisi semuanya, atau file satu-satu):
  - Detail_Data_*TRX.csv   : transaksi (JHK, SALES, SPD, STD, APC, %GM, %OOS ...)
  - Detail_Data_*SLA.csv   : delivery (ONTIME / LATE / JUMLAH_DELIVERY) -> %OTD, %LATE
  - Master MDS.xlsx        : daftar toko mini darkstore (opsional, ada daftar bawaan)
  - Report Daily MDS *.xlsx (hasil periode sebelumnya) boleh ikut di dalam ZIP; diabaikan.
Output: Report Daily MDS <tanggal>.xlsx dengan sheet Summary, Breakdown, data harian/MTD,
Pivot delivery, Data MDS, Data SLA MDS (+ Data Raw / Data SLA penuh bila diminta).
"""
from __future__ import annotations

import io
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.common import recalc_with_libreoffice

DEFAULT_MASTER = [
    ("1GQ7", "CHRISTOPEL MIHING"),
    ("1GF4", "CILIK RIWUT KM 1"),
    ("1GQ3", "DIPANDJAITAN SAMPIT"),
    ("1G5D", "KARET"),
    ("1GE4", "PULANG PISAU"),
    ("1GF5", "TEMANGGUNG TILUNG 2"),
    ("X812", "SUSUKAN"),
    ("CG13", "CILINCING KELAPA"),
    ("1J10", "PERMATA SARI INDAH"),
    ("K471", "CITRA"),
    ("TF45", "GOLF ISLAND"),
    ("KA77", "PURI MANSION"),
    ("TD40", "SUVARNA TERRACE 8"),
    ("X528", "CIBINONG GN SINDUR"),
    ("P304", "TUGU HARUM"),
    ("J566", "TIPAR CAKUNG 1 - 3"),
    ("R236", "AHMAD YANI POLMAN"),
    ("R202", "H.A.MUH. ARSYAD 2"),
    ("R151", "MACCOPA MAROS"),
    ("R184", "MH.THAMRIN BONE"),
    ("R264", "TAMANGAPA RAYA 3"),
    ("R287", "SLTN HSNDDIN PINRANG"),
    ("1D73", "LEBAY HASAN"),
    ("1D74", "PROF. SRI SOEDEWI"),
    ("1P35", "KRIDASANA"),
    ("1P41", "KOMYOS SUDARSO"),
    ("1G65", "KARANG ANYAR"),
    ("W227", "KUTACANE KABANJAHE"),
    ("1D85", "PERINTIS"),
    ("R385", "ANDI RADJA BULUKUMBA"),
    ("1PG2", "ADI SUCIPTO KM 8"),
    ("L510", "SUMBER REJO 2"),
    ("1GM1", "BOEJASIN 2"),
    ("1A08", "PUSKOPKAR"),
    ("L548", "CITRA GARDEN 2"),
    ("1GC1", "PALAM RAYA"),
    ("Q389", "BAKTI SERAGA"),
    ("1AJ4", "AHMAD YANI- DUMAI"),
    ("2D07", "MERAPI SUBUR"),
    ("R446", "MARTADINATA MAMUJU"),
    ("1PH9", "KOTA BARU"),
    ("1GD4", "RAYA BATULICIN"),
    ("1P19", "JEMBATAN KAPUAS II"),
    ("1GC2", "AHMAD YANI KM 12"),
    ("1A0Y", "JENSUDPULAUBENGKALIS"),
    ("UA24", "GEDUNGAN SUMENEP"),
    ("1Y02", "LANSOT"),
    ("2DO6", "SIMP BUKIT LESTARI"),
    ("1DR8", "JENSUD 2 TUNGKAL"),
    ("UD50", "SEMEMI SBY"),
    ("2DR5", "MCDERMOTT"),
    ("JD65", "CITRA INDAH AGAVE"),
]

BULAN = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
         "Agustus", "September", "Oktober", "November", "Desember"]
BULAN_SHORT = ["", "Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul",
               "Agu", "Sept", "Okt", "Nov", "Des"]

AGG_COLS = ["KD_STORE", "NAMA_STORE", "NAMA_BRANCH", "Sum of JHK", "Sum of SALES",
            "Sum of SALES_TAGI", "Average of SPD", "Average of STD",
            "Average of APC", "Average of PERCENT_GM", "Average of PCT_OOS_OFMB"]


# ----------------------------------------------------------------------------
# Load & deteksi file
# ----------------------------------------------------------------------------
def _read_any(src, **kw) -> pd.DataFrame:
    """Baca CSV/XLSX dari path atau file-like (mis. Streamlit UploadedFile)."""
    name = str(getattr(src, "name", src)).lower()
    if name.endswith((".xlsx", ".xls")):
        return pd.read_excel(src, **kw)
    return pd.read_csv(src, **kw)


def _norm_cols(df: pd.DataFrame) -> pd.DataFrame:
    df.columns = [str(c).strip().upper() for c in df.columns]
    return df


def _prep(df: pd.DataFrame) -> pd.DataFrame:
    df = _norm_cols(df)
    df["TANGGAL"] = pd.to_datetime(df["TANGGAL"]).dt.normalize()
    df["KD_STORE"] = df["KD_STORE"].astype(str).str.strip().str.upper()
    return df


def load_raw(src) -> pd.DataFrame:
    """Data TRX (kompatibel dengan versi lama: 1 file Detail_Data)."""
    return _prep(_read_any(src, dtype={"KD_STORE": str, "KD_BRANCH": str}))


def load_master(src=None) -> pd.DataFrame:
    if src is None:
        return pd.DataFrame(DEFAULT_MASTER, columns=["KD_STORE", "NAMA_STORE_MASTER"])
    m = _read_any(src, dtype=str)
    m.columns = [c.strip().upper().replace("_", " ") for c in m.columns]
    m = m.rename(columns={"KD STORE": "KD_STORE", "NAMA STORE": "NAMA_STORE_MASTER"})
    m["KD_STORE"] = m["KD_STORE"].str.strip().str.upper()
    return m.drop_duplicates("KD_STORE")


def _header(path) -> list:
    """Nama kolom saja (cepat) untuk menebak jenis file."""
    try:
        if path.lower().endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True)
            first = next(wb.worksheets[0].iter_rows(min_row=1, max_row=1, values_only=True), ())
            wb.close()
            return [str(c).strip().upper() for c in first if c is not None]
        return [c.strip().upper() for c in pd.read_csv(path, nrows=0).columns]
    except Exception:  # noqa: BLE001
        return []


def classify(paths) -> dict:
    """Kelompokkan file: trx, sla, master, previous (report lama), other."""
    out = {"trx": None, "sla": None, "master": None, "previous": None, "other": []}
    for p in paths:
        base = os.path.basename(p).lower()
        if not base.endswith((".csv", ".xlsx", ".xls", ".xlsm")):
            continue
        if base.startswith("report daily mds") or base.startswith("report_daily_mds"):
            out["previous"] = p            # hasil laporan lama: tidak dibaca (besar), hanya dikenali
            continue
        cols = _header(p)
        if "DELIVERY_ONTIME" in cols or "JUMLAH_DELIVERY" in cols:
            out["sla"] = out["sla"] or p
        elif "JHK" in cols and "SALES" in cols:
            out["trx"] = out["trx"] or p
        elif {"KD STORE", "KD_STORE"} & set(cols) and "TANGGAL" not in cols:
            out["master"] = out["master"] or p
        elif "master" in base:
            out["master"] = out["master"] or p
        else:
            out["other"].append(p)
    return out


def load_inputs(paths, master_path=None) -> dict:
    """Baca TRX + SLA + master dari daftar path. SLA boleh tidak ada (kolom %OTD kosong)."""
    g = classify(paths)
    if g["trx"] is None:
        raise FileNotFoundError(
            "File TRX (Detail_Data_*TRX.csv, kolom JHK/SALES) tidak ditemukan di upload.")
    trx = load_raw(g["trx"])
    sla = _prep(_read_any(g["sla"], dtype={"KD_STORE": str, "KD_BRANCH": str})) if g["sla"] else None
    mp = master_path or g["master"]
    master = load_master(mp)
    return dict(trx=trx, sla=sla, master=master, files=g, master_path=mp)


def available_dates(src):
    """Tanggal yang bisa dipilih. src = path TRX (lama) atau dict hasil load_inputs()."""
    if isinstance(src, dict):
        d = set(src["trx"]["TANGGAL"].dt.date.unique())
        if src["sla"] is not None:
            d &= set(src["sla"]["TANGGAL"].dt.date.unique())
        return sorted(d)
    return sorted(load_raw(src)["TANGGAL"].dt.date.unique())


# ----------------------------------------------------------------------------
# Transform
# ----------------------------------------------------------------------------
def _agg(df: pd.DataFrame) -> pd.DataFrame:
    g = df.groupby(["KD_STORE", "NAMA_STORE", "NAMA_BRANCH"], as_index=False).agg(
        **{
            "Sum of JHK": ("JHK", "sum"),
            "Sum of SALES": ("SALES", "sum"),
            "Sum of SALES_TAGI": ("SALES_TAGI", "sum"),
            "Average of SPD": ("SPD", "mean"),
            "Average of STD": ("STD", "mean"),
            "Average of APC": ("APC", "mean"),
            "Average of PERCENT_GM": ("PERCENT_GM", "mean"),
            "Average of PCT_OOS_OFMB": ("PCT_OOS_OFMB", "mean"),
        }
    )
    return g[AGG_COLS]


def _agg_sla(df: pd.DataFrame) -> pd.DataFrame:
    g = df.groupby(["KD_STORE", "NAMA_STORE", "NAMA_BRANCH"], as_index=False).agg(
        **{
            "Sum of DELIVERY_ONTIME": ("DELIVERY_ONTIME", "sum"),
            "Sum of DELIVERY_LATE": ("DELIVERY_LATE", "sum"),
            "Sum of JUMLAH_DELIVERY": ("JUMLAH_DELIVERY", "sum"),
        })
    g["%OTD"] = g["Sum of DELIVERY_ONTIME"] / g["Sum of JUMLAH_DELIVERY"]
    g["%LATE"] = g["Sum of DELIVERY_LATE"] / g["Sum of JUMLAH_DELIVERY"]
    return g


def _tag_mds(df, master):
    df = df.copy()
    ins = df["KD_STORE"].isin(master["KD_STORE"])
    pos = df.columns.get_loc("REMARK") + 1 if "REMARK" in df.columns else len(df.columns)
    df.insert(pos, "REMARK MDS", ins.map({True: "MINI DARKSTORE", False: "TIDAK"}))
    return df


def build(raw: pd.DataFrame, master: pd.DataFrame, tanggal=None, sla: pd.DataFrame | None = None) -> dict:
    if "PCT_OOS_OFMB" not in raw.columns:
        raw = raw.assign(PCT_OOS_OFMB=float("nan"))
    tgl = pd.Timestamp(tanggal).normalize() if tanggal else raw["TANGGAL"].max()
    if tanggal is None and sla is not None:
        tgl = min(tgl, sla["TANGGAL"].max())
    awal = tgl.replace(day=1)

    raw_t = _tag_mds(raw, master)
    data_mds = raw_t[raw_t["REMARK MDS"] == "MINI DARKSTORE"].copy()
    mtd_rows = data_mds[(data_mds["TANGGAL"] >= awal) & (data_mds["TANGGAL"] <= tgl)]
    day_rows = data_mds[data_mds["TANGGAL"] == tgl]
    if day_rows.empty:
        raise ValueError(f"Tidak ada data MDS untuk tanggal {tgl:%Y-%m-%d}")

    daily, mtd = _agg(day_rows), _agg(mtd_rows)

    sla_raw = sla_mds = sla_day = sla_mtd = None
    if sla is not None:
        sla_raw = _tag_mds(sla, master)
        sla_mds = sla_raw[sla_raw["REMARK MDS"] == "MINI DARKSTORE"].copy()
        sla_day = _agg_sla(sla_mds[sla_mds["TANGGAL"] == tgl])
        sla_mtd = _agg_sla(sla_mds[(sla_mds["TANGGAL"] >= awal) & (sla_mds["TANGGAL"] <= tgl)])

    s = mtd.merge(daily, on=["KD_STORE"], how="left", suffixes=("_mtd", "_d"))
    if sla is not None:
        s = s.merge(sla_day[["KD_STORE", "%OTD", "%LATE"]].rename(columns={"%OTD": "OTD_d", "%LATE": "LATE_d"}),
                    on="KD_STORE", how="left")
        s = s.merge(sla_mtd[["KD_STORE", "%OTD", "%LATE"]].rename(columns={"%OTD": "OTD_m", "%LATE": "LATE_m"}),
                    on="KD_STORE", how="left")
    else:
        for c in ("OTD_d", "LATE_d", "OTD_m", "LATE_m"):
            s[c] = float("nan")

    summary = pd.DataFrame({
        "Kode Toko": s["KD_STORE"],
        "Nama Toko": s["NAMA_STORE_mtd"],
        "Cabang": s["NAMA_BRANCH_mtd"],
        "D_JHK": s["Sum of JHK_d"],
        "D_SPD": s["Average of SPD_d"],
        "D_SPD_TAGI": s["Sum of SALES_TAGI_d"],
        "D_STD": s["Average of STD_d"],
        "D_APC": s["Average of APC_d"],
        "D_GM": s["Average of PERCENT_GM_d"],
        "D_OOS": s["Average of PCT_OOS_OFMB_d"],
        "D_OTD": s["OTD_d"],
        "D_LATE": s["LATE_d"],
        "M_JHK": s["Sum of JHK_mtd"],
        "M_SALES": s["Sum of SALES_mtd"],
        "M_SALES_TAGI": s["Sum of SALES_TAGI_mtd"],
        "M_SPD": s["Average of SPD_mtd"],
        "M_SPD_TAGI": s["Sum of SALES_TAGI_mtd"] / s["Sum of JHK_mtd"],
        "M_STD": s["Average of STD_mtd"],
        "M_APC": s["Average of APC_mtd"],
        "M_GM": s["Average of PERCENT_GM_mtd"],
        "M_OOS": s["Average of PCT_OOS_OFMB_mtd"],
        "M_OTD": s["OTD_m"],
        "M_LATE": s["LATE_m"],
    }).sort_values("D_SPD", ascending=False, ignore_index=True)   # urutan sama dengan report manual

    missing = sorted(set(master["KD_STORE"]) - set(data_mds["KD_STORE"]))
    return dict(tgl=tgl, awal=awal, summary=summary, daily=daily, mtd=mtd,
                data_mds=data_mds.sort_values(["TANGGAL", "KD_STORE"]),
                raw=raw_t, master=master, sla_raw=sla_raw,
                sla_mds=None if sla_mds is None else sla_mds.sort_values(["TANGGAL", "KD_STORE"]),
                sla_day=sla_day, sla_mtd=sla_mtd, missing=missing)


# ----------------------------------------------------------------------------
# Write Excel
# ----------------------------------------------------------------------------
FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)


def _clean(v):
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if hasattr(v, "item") and not isinstance(v, (str, bytes)):
        try:
            return v.item()
        except Exception:  # noqa: BLE001
            return v
    return v


def _write_table(ws, df: pd.DataFrame, num_fmt: dict | None = None, widths=True):
    ws.append(list(df.columns))
    for c in ws[1]:
        c.font, c.fill = BOLD, FILL
    for row in df.itertuples(index=False):
        ws.append([_clean(v) for v in row])
    for i, col in enumerate(df.columns, 1):
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(28, len(str(col)) + 2))
        if num_fmt and col in num_fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = num_fmt[col]
    ws.freeze_panes = "A2"


def _summary_sheet(wb, res, d_label, p_label):
    sm = res["summary"]
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws["B2"] = "Daily Report Performance Mini Darkstore"
    ws["B3"] = f"Periode {p_label}"
    ws["B2"].font = TITLE
    ws["B3"].font = BOLD

    for col, txt in zip("BCDE", ["No", "Kode Toko", "Nama Toko", "Cabang"]):
        ws[f"{col}5"] = txt
        ws.merge_cells(f"{col}5:{col}6")
    ws["F5"] = f"Sales Mini Darkstore {d_label}"
    ws.merge_cells("F5:N5")
    ws["O5"] = f"Sales Mini Darkstore {p_label}"
    ws.merge_cells("O5:Y5")
    sub = ["JHK", "SPD", "SPD Tag I", "STD", "APC", "%GM", "%OOS OFMB", "%OTD", "%LATE",
           "JHK", "Total Net Sales", "Total Net Sales Tag I", "SPD", "SPD Tag I", "STD", "APC",
           "%GM", "%OOS OFMB", "%OTD", "%LATE"]
    for i, txt in enumerate(sub):
        ws.cell(6, 6 + i, txt)
    for r in (5, 6):
        for c in range(2, 26):
            cell = ws.cell(r, c)
            cell.font, cell.fill, cell.alignment, cell.border = BOLD, FILL, CENTER, BORDER

    keys = ["D_JHK", "D_SPD", "D_SPD_TAGI", "D_STD", "D_APC", "D_GM", "D_OOS", "D_OTD", "D_LATE",
            "M_JHK", "M_SALES", "M_SALES_TAGI", "M_SPD", "M_SPD_TAGI", "M_STD", "M_APC", "M_GM",
            "M_OOS", "M_OTD", "M_LATE"]
    fmts = ["0", "#,##0", "#,##0", "#,##0", "#,##0", "0.00%", "0.00%", "0.00%", "0.00%",
            "0", "#,##0", "#,##0", "#,##0", "#,##0", "#,##0", "#,##0", "0.00%", "0.00%", "0.00%", "0.00%"]
    first = 7
    for n, row in sm.iterrows():
        r = first + n
        vals = [n + 1, row["Kode Toko"], row["Nama Toko"], row["Cabang"]] + [row[k] for k in keys]
        for c, v in enumerate(vals, 2):
            cell = ws.cell(r, c, _clean(v))
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="left" if c in (4, 5) else "center")
            if c >= 6:
                cell.number_format = fmts[c - 6]
    last = first + len(sm) - 1

    tr = last + 1
    ws.cell(tr, 2, "Average/Sum")
    ws.merge_cells(start_row=tr, start_column=2, end_row=tr, end_column=5)
    for c in range(6, 26):
        L = get_column_letter(c)
        fn = "SUM" if L in ("P", "Q") else "AVERAGE"
        ws.cell(tr, c, f"={fn}({L}{first}:{L}{last})").number_format = fmts[c - 6]
    for c in range(2, 26):
        cell = ws.cell(tr, c)
        cell.font, cell.fill, cell.alignment, cell.border = BOLD, FILL, CENTER, BORDER

    ws.conditional_formatting.add(
        f"P{first}:P{last}",
        ColorScaleRule(start_type="min", start_color="FFEF9C", end_type="max", end_color="63BE7B"))
    ws.conditional_formatting.add(
        f"X{first}:X{last}",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50,
                       mid_color="FFEB84", end_type="max", end_color="63BE7B"))

    widths = dict(A=2, B=5, C=10, D=26, E=14, F=6, G=13, H=13, I=7, J=10, K=8, L=10, M=8, N=8,
                  O=6, P=16, Q=18, R=13, S=12, T=7, U=10, V=8, W=10, X=8, Y=8)
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.row_dimensions[5].height = 30
    ws.freeze_panes = "F7"


def _breakdown_sheet(wb, res, p_label):
    sm = res["summary"]
    ws = wb.create_sheet("Breakdown")
    ws.sheet_view.showGridLines = False
    ws["B2"] = "Performance Analysis Mini Darkstore - SPD & %OTD"
    ws["B3"] = f"Periode {p_label}"
    ws["B2"].font = TITLE
    ws["B3"].font = BOLD
    for c, t in zip("BCDEFGH", ["Kode Toko", "Nama Toko", "Cabang", "SPD", "%OTD", "Kategori SPD", "Kategori %OTD"]):
        ws[f"{c}5"] = t
    for c in "BCDEFGH":
        ws[f"{c}5"].font, ws[f"{c}5"].fill, ws[f"{c}5"].alignment, ws[f"{c}5"].border = BOLD, FILL, CENTER, BORDER

    b = sm.sort_values("M_OTD", ascending=False, na_position="last", ignore_index=True) if False else sm
    r0 = 6
    r1 = r0 + len(b) - 1
    for i, row in b.iterrows():
        r = r0 + i
        ws.cell(r, 2, row["Kode Toko"]); ws.cell(r, 3, row["Nama Toko"]); ws.cell(r, 4, row["Cabang"])
        ws.cell(r, 5, _clean(row["M_SPD"])).number_format = "#,##0"
        ws.cell(r, 6, _clean(row["M_OTD"])).number_format = "0.0%"
        ws.cell(r, 7, f'=IF(E{r}<$K$6,"Bottom 25%",IF(E{r}<$L$6,"25-50%",IF(E{r}<$M$6,"50-75%","Top 25%")))')
        ws.cell(r, 8, f'=IF(F{r}<$K$7,"Bottom 25%",IF(F{r}<$L$7,"25-50%",IF(F{r}<$M$7,"50-75%","Top 25%")))')
        for c in range(2, 9):
            ws.cell(r, c).border = BORDER

    for c, t in zip("JKLM", ["Metrik", "Q1", "Q2", "Q3"]):
        ws[f"{c}5"] = t
        ws[f"{c}5"].font, ws[f"{c}5"].fill, ws[f"{c}5"].alignment, ws[f"{c}5"].border = BOLD, FILL, CENTER, BORDER
    ws["J6"], ws["J7"] = "SPD", "%OTD"
    for q, col in zip((1, 2, 3), "KLM"):
        ws[f"{col}6"] = f"=_xlfn.QUARTILE.INC(E{r0}:E{r1},{q})"
        ws[f"{col}7"] = f"=_xlfn.QUARTILE.INC(F{r0}:F{r1},{q})"
        ws[f"{col}6"].number_format = "#,##0"
        ws[f"{col}7"].number_format = "0.0%"

    ws["J9"] = "Matriks SPD vs %OTD"
    ws["J9"].font = BOLD
    ws["J10"], ws["K10"] = "SPD", "%OTD"
    cats = ["Bottom 25%", "25-50%", "50-75%", "Top 25%"]
    for j, cname in enumerate(cats):
        ws.cell(11, 11 + j, cname)
        ws.cell(12 + j, 10, cname)
    for c in range(10, 15):
        for r in (10, 11):
            ws.cell(r, c).font, ws.cell(r, c).fill, ws.cell(r, c).alignment = BOLD, FILL, CENTER
    for i in range(4):
        r = 12 + i
        ws.cell(r, 10).font = BOLD
        for j in range(4):
            col = get_column_letter(11 + j)
            ws.cell(r, 11 + j, f"=COUNTIFS($G${r0}:$G${r1},$J{r},$H${r0}:$H${r1},{col}$11)")
            ws.cell(r, 11 + j).border = BORDER
            ws.cell(r, 11 + j).alignment = Alignment(horizontal="center")
    ws.conditional_formatting.add(
        "K12:N15", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="63BE7B"))

    for k, v in dict(A=2, B=10, C=26, D=14, E=14, F=9, G=14, H=14, I=3, J=18, K=14, L=14, M=14, N=12).items():
        ws.column_dimensions[k].width = v
    ws.freeze_panes = "A6"


def to_excel(res: dict, include_raw: bool = False) -> bytes:
    tgl, awal = res["tgl"], res["awal"]
    bln, thn = BULAN[tgl.month], tgl.year
    d_label = f"{tgl.day} {bln} {thn}"
    p_label = f"{awal.day} - {tgl.day} {bln} {thn}"
    sh_day = f"Data {tgl.day} {BULAN_SHORT[tgl.month]} {thn}"
    sh_mtd = f"Data MTD {awal.day} - {tgl.day} {BULAN_SHORT[tgl.month]} {thn}"

    wb = Workbook()
    _summary_sheet(wb, res, d_label, p_label)
    _breakdown_sheet(wb, res, p_label)

    pfmt = {"Sum of SALES": "#,##0", "Sum of SALES_TAGI": "#,##0", "Average of SPD": "#,##0",
            "Average of STD": "#,##0.0", "Average of APC": "#,##0", "Average of PERCENT_GM": "0.00%",
            "Average of PCT_OOS_OFMB": "0.00%"}
    _write_table(wb.create_sheet(sh_day[:31]), res["daily"], pfmt)
    _write_table(wb.create_sheet(sh_mtd[:31]), res["mtd"], pfmt)

    if res["sla_day"] is not None:
        sfmt = {"%OTD": "0.00%", "%LATE": "0.00%"}
        _write_table(wb.create_sheet(f"Delivery {tgl.day} {BULAN_SHORT[tgl.month]}"[:31]), res["sla_day"], sfmt)
        _write_table(wb.create_sheet(f"Delivery MTD {awal.day} - {tgl.day} {BULAN_SHORT[tgl.month]}"[:31]),
                     res["sla_mtd"], sfmt)

    _write_table(wb.create_sheet("Master MDS"), res["master"].rename(
        columns={"KD_STORE": "KD STORE", "NAMA_STORE_MASTER": "NAMA STORE"}))
    _write_table(wb.create_sheet("Data MDS"), res["data_mds"], {"TANGGAL": "yyyy-mm-dd"}, widths=False)
    if res["sla_mds"] is not None:
        _write_table(wb.create_sheet("Data SLA MDS"), res["sla_mds"], {"TANGGAL": "yyyy-mm-dd"}, widths=False)
    if include_raw:
        _write_table(wb.create_sheet("Data Raw"), res["raw"], {"TANGGAL": "yyyy-mm-dd"}, widths=False)
        if res["sla_raw"] is not None:
            _write_table(wb.create_sheet("Data SLA"), res["sla_raw"], {"TANGGAL": "yyyy-mm-dd"}, widths=False)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate(raw_src, master_src=None, tanggal=None, sla_src=None, include_raw=False) -> tuple[bytes, dict]:
    sla = _prep(_read_any(sla_src, dtype={"KD_STORE": str, "KD_BRANCH": str})) if sla_src is not None else None
    res = build(load_raw(raw_src), load_master(master_src), tanggal, sla)
    return to_excel(res, include_raw), res


def output_name(res: dict) -> str:
    t = res["tgl"]
    return f"Report Daily MDS {t.day} {BULAN[t.month]} {t.year}.xlsx"


def run(paths, output_dir, master_path=None, tanggal=None, include_raw=False, log=print):
    """paths = daftar file hasil ekstrak (TRX + SLA + Master [+ report lama]); ZIP sudah diekstrak
    oleh prepare_inputs()."""
    inp = load_inputs(paths, master_path)
    g = inp["files"]
    log("File TRX     :", os.path.basename(g["trx"]))
    log("File SLA     :", os.path.basename(g["sla"]) if g["sla"] else "TIDAK ADA -> kolom %OTD/%LATE kosong")
    mp = inp["master_path"]
    log("Master MDS   :", os.path.basename(mp) if mp else f"bawaan ({len(DEFAULT_MASTER)} toko)")
    if g["previous"]:
        log("Report lama  :", os.path.basename(g["previous"]), "(dikenali, tidak dipakai sebagai input)")
    if tanggal is None:
        dates = available_dates(inp)
        if not dates:
            raise ValueError("Tidak ada tanggal yang sama antara file TRX dan SLA.")
        tanggal = dates[-1]
    res = build(inp["trx"], inp["master"], tanggal, inp["sla"])
    from core import storage  # simpan snapshot MDS ke Google Sheet (tab 'mds')
    storage.save_mds_if_configured(res, log)
    xlsx = to_excel(res, include_raw)
    sm = res["summary"]
    log(f"Tanggal report: {res['tgl']:%d-%m-%Y} | {len(sm)} toko | "
        f"Total Net Sales MTD Rp {sm['M_SALES'].sum():,.0f}")
    if res["missing"]:
        log(f"[Peringatan] Toko di master tapi tidak ada di data: {', '.join(res['missing'])}")
    out = os.path.join(output_dir, output_name(res))
    with open(out, "wb") as f:
        f.write(xlsx)
    recalc_with_libreoffice(out, log)
    preview = sm.rename(columns={
        "D_JHK": "JHK", "D_SPD": "SPD", "D_SPD_TAGI": "SPD Tag I", "D_STD": "STD", "D_APC": "APC",
        "D_GM": "%GM", "D_OOS": "%OOS OFMB", "D_OTD": "%OTD", "D_LATE": "%LATE",
        "M_JHK": "JHK MTD", "M_SALES": "Net Sales MTD", "M_SALES_TAGI": "Net Sales Tag I MTD",
        "M_SPD": "SPD MTD", "M_SPD_TAGI": "SPD Tag I MTD", "M_STD": "STD MTD", "M_APC": "APC MTD",
        "M_GM": "%GM MTD", "M_OOS": "%OOS OFMB MTD", "M_OTD": "%OTD MTD", "M_LATE": "%LATE MTD"})
    return out, preview, []
