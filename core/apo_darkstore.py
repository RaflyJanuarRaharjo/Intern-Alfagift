"""Update APO Darkstore.

Sumber: notebooks/master-data/update_apo_darkstore_bitmap.ipynb.
Input tanggal/jam yang dulu lewat input() kini lewat form di web.
Tambahan: tabel juga di-render jadi gambar PNG (siap kirim ke grup chat),
sesuai tujuan notebook ("hasil sebagai bitmap").
"""
import numbers
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .common import Log

URUTAN_STATUS = ["1. New", "2. Packing", "3. Pesanan Siap", "4. Siap Kirim",
                 "5. Dalam Pengiriman", "6. Tunda", "7. Selesai", "8. Batal"]
HEADER_COLOR = "B7DEE8"
LEBAR = [7, 34, 10, 12, 15, 13, 20, 10, 11, 10, 13]  # B..L


def load_data(csv_path):
    data = pd.read_csv(csv_path)
    data.columns = data.columns.str.strip()
    data = data.rename(columns={"store": "Nama Toko", "sort status apo": "Status"})
    for col in ("Nama Toko", "Status", "Jumlah Order"):
        if col not in data.columns:
            raise KeyError(f"Kolom '{col}' tidak ada. Kolom yang tersedia: {list(data.columns)}")
    data["Jumlah Order"] = pd.to_numeric(data["Jumlah Order"], errors="coerce").fillna(0)
    data["Nama Toko"] = data["Nama Toko"].astype(str).str.strip()
    data["Status"] = data["Status"].astype(str).str.strip()
    return data


def build_table(data):
    tabel = pd.pivot_table(data, index="Nama Toko", columns="Status", values="Jumlah Order",
                           aggfunc="sum", fill_value=0)
    tabel = tabel.reindex(columns=URUTAN_STATUS, fill_value=0)
    tabel["Grand Total"] = tabel.sum(axis=1)
    tabel = tabel.sort_values(by="7. Selesai", ascending=True).reset_index()
    tabel.columns.name = None
    tabel.insert(0, "No", range(1, len(tabel) + 1))
    total = {"No": "", "Nama Toko": "Grand Total"}
    for kolom in URUTAN_STATUS + ["Grand Total"]:
        total[kolom] = tabel[kolom].sum()
    return pd.concat([tabel, pd.DataFrame([total])], ignore_index=True)


def selesai_colors(n_toko):
    """Gradasi merah -> kuning -> hijau untuk kolom 7. Selesai (hex RRGGBB)."""
    out = []
    for i in range(n_toko):
        pos = i / (n_toko - 1) if n_toko > 1 else 0
        if pos <= 0.5:
            merah, hijau = 255, int(255 * (pos / 0.5))
        else:
            merah, hijau = int(255 * (1 - (pos - 0.5) / 0.5)), 255
        out.append(f"{merah:02X}{hijau:02X}00")
    return out


def write_excel(tabel, judul_waktu, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "UPDATE APO DARKSTORE"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("B1:L1")
    ws["B1"] = "UPDATE APO DARKSTORE"
    ws["B1"].font = Font(bold=True, size=14)
    ws["B1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("B2:L2")
    ws["B2"] = judul_waktu
    ws["B2"].font = Font(bold=True, size=11)
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center")

    garis = Side(style="thin", color="000000")
    border = Border(left=garis, right=garis, top=garis, bottom=garis)
    center = Alignment(horizontal="center", vertical="center")

    for kolom, nama in enumerate(tabel.columns, 2):
        cell = ws.cell(row=4, column=kolom, value=nama)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor=HEADER_COLOR)
        cell.alignment = center
        cell.border = border

    for b in range(len(tabel)):
        for k in range(len(tabel.columns)):
            val = tabel.iloc[b, k]
            if hasattr(val, "item"):
                val = val.item()
            cell = ws.cell(row=b + 5, column=k + 2, value=val)
            cell.border = border
            cell.alignment = Alignment(horizontal="left", vertical="center") if k == 1 else center

    for i, warna in enumerate(selesai_colors(len(tabel) - 1)):
        ws.cell(row=i + 5, column=10).fill = PatternFill(fill_type="solid", fgColor=warna)

    baris_total = len(tabel) + 4
    ws.merge_cells(start_row=baris_total, start_column=2, end_row=baris_total, end_column=3)
    ws.cell(row=baris_total, column=2).value = "Grand Total"
    for kolom in range(2, 13):
        cell = ws.cell(row=baris_total, column=kolom)
        cell.fill = PatternFill(fill_type="solid", fgColor=HEADER_COLOR)
        cell.font = Font(bold=True)
        cell.alignment = center
        cell.border = border

    ws.column_dimensions["A"].width = 8.4
    for kolom, ukuran in enumerate(LEBAR, 2):
        ws.column_dimensions[get_column_letter(kolom)].width = ukuran
    ws.column_dimensions["M"].width = 8.4
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = "B5"
    wb.save(path)


def write_png(tabel, judul_waktu, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    n_rows = len(tabel)
    fig_h = 0.7 + 0.24 * (n_rows + 1)
    fig, ax = plt.subplots(figsize=(13, fig_h), dpi=150)
    ax.axis("off")
    def fmt(k, v):
        if k > 1 and isinstance(v, numbers.Number):
            return f"{int(v):,}".replace(",", ".")
        return str(v)

    cells = [[fmt(k, v) for k, v in enumerate(row)] for row in tabel.itertuples(index=False)]
    t = ax.table(cellText=cells, colLabels=list(tabel.columns), loc="upper center", cellLoc="center",
                 colWidths=[w / sum(LEBAR) for w in LEBAR])
    t.auto_set_font_size(False)
    t.set_fontsize(8.5)
    t.scale(1, 1.35)
    colors = selesai_colors(n_rows - 1)
    for (r, c), cell in t.get_celld().items():
        cell.set_edgecolor("#000000")
        cell.set_linewidth(0.5)
        if r == 0 or r == n_rows:
            cell.set_facecolor("#" + HEADER_COLOR)
            cell.set_text_props(weight="bold")
        elif c == 8:
            cell.set_facecolor("#" + colors[r - 1])
        if c == 1 and r > 0:
            cell.set_text_props(ha="left")
            cell._loc = "left"
    ax.set_title(f"UPDATE APO DARKSTORE\n{judul_waktu}", loc="left", fontsize=12, fontweight="bold")
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def run(paths, output_dir, tanggal, jam, make_png=True, log=None):
    log = log or Log()
    csvs = [p for p in paths if p.lower().endswith(".csv")]
    if not csvs:
        raise FileNotFoundError("File CSV (GLI - NEW SAPA Performance ... Tabel pivot.csv) tidak ditemukan.")
    log("File yang digunakan:", os.path.basename(csvs[0]))
    tabel = build_table(load_data(csvs[0]))
    judul_waktu = f"{tanggal} JAM {jam}"
    out = os.path.join(output_dir, "UPDATE_APO_DARKSTORE.xlsx")
    write_excel(tabel, judul_waktu, out)
    log(f"{len(tabel) - 1} toko, diurutkan dari '7. Selesai' terkecil.")
    extra = []
    if make_png:
        png = os.path.join(output_dir, "UPDATE_APO_DARKSTORE.png")
        write_png(tabel, judul_waktu, png)
        extra.append(png)
    return out, tabel, extra
