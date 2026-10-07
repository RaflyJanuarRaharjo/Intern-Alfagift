"""Report Inventory, OOS, MAT, Store Performance.

Sumber: notebooks/daily/report_daily_inventory.ipynb. Logika pipeline
dipindah apa adanya; template bawaan kini disimpan sebagai file di
templates/default_inventory_template.xlsx (sebelumnya base64 di notebook).
"""
import datetime as dt
import gc
import os
import re
from copy import copy

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.utils import get_column_letter

from .common import Log

MONTH_ID = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
            "Agustus", "September", "Oktober", "November", "Desember"]

DEFAULT_TEMPLATE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                "templates", "default_inventory_template.xlsx")

REQUIRED_KEYS = ["summary_mtd", "oos", "list_ds", "inventory", "mat_all", "mat_exc"]
LABELS = {
    "template": "Report periode sebelumnya (template, opsional)",
    "summary_mtd": "Summary_MTD_*.xlsx",
    "oos": "[OOS] By Toko.xlsx",
    "list_ds": "List DS_Covered Invent_*.xlsx",
    "inventory": "inventory_darkstore_*.xlsx",
    "mat_all": "[MAT] Transaksi All.xlsx",
    "mat_exc": "[MAT] Transaksi Excl Beanspot.xlsx",
}


# ---------------------------------------------------------------- loaders
def _rows(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    wb.close()
    del wb
    gc.collect()
    return rows


def load_kd_store(path):
    return [r for r in _rows(path, 'KD_STORE') if any(v is not None for v in r)]


def load_oos(path):
    rows = [r for r in _rows(path, 'Sheet 1') if any(v is not None for v in r)]
    return [r for r in rows if str(r[0]).strip().lower() != 'grand total']


def load_list_ds(path):
    return [r[:6] for r in _rows(path, 'List DS') if r[0] is not None]


def load_inventory(path):
    all_rows = _rows(path, 'Store')
    period_text = all_rows[1][0]  # 'PERIODE 01-09-2026 SD 07-09-2026'
    header_idx = next(i for i, r in enumerate(all_rows) if r[0] == 'Branch Code')
    header = all_rows[header_idx]
    data = all_rows[header_idx + 1:]
    remark_idx = header.index('Remark Item')
    branch_name_idx = header.index('Branch Name')
    filtered = [r for r in data if r[0] is not None and r[remark_idx] == 'ALL']
    keep_idx = [i for i in range(len(header)) if i not in (0, branch_name_idx, remark_idx)]
    out = [['No'] + [header[i] for i in keep_idx]]
    for n, r in enumerate(filtered, start=1):
        out.append([n] + [r[i] for i in keep_idx])
    m = re.search(r'(\d{2})-(\d{2})-(\d{4})\s+SD\s+(\d{2})-(\d{2})-(\d{4})', str(period_text))
    if not m:
        raise ValueError(f"Format periode di file Inventory tidak dikenali: {period_text!r}")
    start = dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    end = dt.date(int(m.group(6)), int(m.group(5)), int(m.group(4)))
    return out, start, end


def load_mat_long(path, value_label):
    rows = _rows(path, 'Sheet 1')
    codes, names = rows[0][1:], rows[1][1:]
    avg_row = next(r[1:] for r in rows[2:] if r[0] == 'Average')
    out = [['No', 'Kode Toko', 'Nama Toko', value_label]]
    for i, (code, name, val) in enumerate(zip(codes, names, avg_row), start=1):
        out.append([i, code, name, val])
    return out


# ---------------------------------------------------------------- writers
def _copy_row_style(ws, src_row, dst_row, ncols):
    for c in range(1, ncols + 1):
        s, d = ws.cell(row=src_row, column=c), ws.cell(row=dst_row, column=c)
        d.font, d.fill, d.border = copy(s.font), copy(s.fill), copy(s.border)
        d.alignment, d.number_format = copy(s.alignment), s.number_format


def write_table_sheet(wb, sheet_name, table_name, rows):
    """Timpa area tabel; tidak pakai insert_rows/delete_rows (lambat di sheet besar)."""
    ws = wb[sheet_name]
    table = ws.tables[table_name]
    min_col, min_row, max_col, max_row = openpyxl.utils.cell.range_boundaries(table.ref)
    old_n, new_n, ncols = max_row - min_row, len(rows) - 1, len(rows[0])
    for r_off, row in enumerate(rows):
        for c_off, val in enumerate(row):
            ws.cell(row=min_row + r_off, column=min_col + c_off, value=val)
    new_last = min_row + new_n
    if new_n > old_n:
        src = max_row if old_n > 0 else min_row + 1
        for r in range(max_row + 1, new_last + 1):
            _copy_row_style(ws, src, r, ncols)
    elif new_n < old_n:
        for r in range(new_last + 1, max_row + 1):
            for c in range(min_col, max_col + 1):
                ws.cell(row=r, column=c, value=None)
    table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{new_last}"


REPORT_FORMULA_COLS = {
    5: "=VLOOKUP(C{r},Table8[],6,FALSE)",
    6: "=VLOOKUP(C{r},Table8[],7,FALSE)",
    7: "=VLOOKUP(C{r},Table8[],11,FALSE)",
    8: "=VLOOKUP(C{r},Table8[],12,FALSE)",
    9: "=VLOOKUP(C{r},Table8[],8,FALSE)",
    10: "=VLOOKUP(C{r},Table25[[Kode Toko]:[MAT All]],3,FALSE)",
    11: "=VLOOKUP(C{r},Table2[[#All],[Kode Toko]:[MAT Exc Beanspot]],3,FALSE)",
    12: "=VLOOKUP(C{r},Table6[],4,FALSE)",
    13: "=VLOOKUP(C{r},Table6[],5,FALSE)",
    14: "=VLOOKUP(C{r},Table5[[Store Code]:[MAT % Beanspot]],8,FALSE)",
    15: "=VLOOKUP(C{r},Table5[[Store Code]:[MAT % Beanspot]],10,FALSE)",
    16: "=VLOOKUP(C{r},Table5[[Store Code]:[MAT % Beanspot]],13,FALSE)",
    17: "=VLOOKUP(C{r},Table5[[Store Code]:[MAT % Beanspot]],15,FALSE)",
    18: "=VLOOKUP(C{r},Table1[],6,FALSE)",
}


def _shift_conditional_formatting(ws, old_first, old_last, new_first, new_last):
    new_cf = ConditionalFormattingList()
    for cf in ws.conditional_formatting:
        new_ranges = []
        for rng in str(cf.sqref).split():
            m = re.match(r'^([A-Z]+)(\d+):([A-Z]+)(\d+)$', rng)
            if m and int(m.group(2)) == old_first and int(m.group(4)) == old_last:
                new_ranges.append(f"{m.group(1)}{new_first}:{m.group(3)}{new_last}")
            else:
                new_ranges.append(rng)
        for rule in cf.rules:
            new_cf.add(" ".join(new_ranges), rule)
    ws.conditional_formatting = new_cf


def rebuild_report_sheet(wb, kd_store_rows, list_ds_rows, inventory_rows, period_start, period_end):
    ws = wb['Report']
    kd_header = kd_store_rows[0]
    i_kode, i_branch, i_name = (kd_header.index(k) for k in ('KD_STORE', 'NAMA_BRANCH', 'NAMA_STORE'))
    kd_map = {r[i_kode]: (r[i_branch], r[i_name]) for r in kd_store_rows[1:]}

    inv_header = inventory_rows[0]
    i_code, i_val = inv_header.index('Store Code'), inv_header.index('Inventory Value')
    inv_value_map = {r[i_code]: r[i_val] for r in inventory_rows[1:]}

    store_codes = [r[0] for r in list_ds_rows[1:]]
    store_codes.sort(key=lambda code: inv_value_map.get(code) or 0, reverse=True)

    OLD_FIRST, OLD_LAST = 6, 72
    old_n, new_n = OLD_LAST - OLD_FIRST + 1, len(store_codes)
    new_last = OLD_FIRST + new_n - 1
    for i, code in enumerate(store_codes):
        r = OLD_FIRST + i
        branch, name = kd_map.get(code, (None, None))
        ws.cell(row=r, column=2, value=branch)
        ws.cell(row=r, column=3, value=code)
        ws.cell(row=r, column=4, value=name)
        for col, tmpl in REPORT_FORMULA_COLS.items():
            ws.cell(row=r, column=col, value=tmpl.format(r=r))
        if new_n > old_n and r > OLD_LAST:
            _copy_row_style(ws, OLD_LAST, r, 18)
    if new_n < old_n:
        for r in range(new_last + 1, OLD_LAST + 1):
            for c in range(2, 19):
                ws.cell(row=r, column=c, value=None)
    if new_n != old_n:
        _shift_conditional_formatting(ws, OLD_FIRST, OLD_LAST, OLD_FIRST, new_last)

    bulan = MONTH_ID[period_end.month - 1]
    rng = f"{period_start.day} - {period_end.day}"
    ws['B3'] = f"Report Inventory MTD {rng} {bulan} {period_end.year}"
    ws['J4'] = f"MAT (%) [{rng} {bulan[:3]}]"
    ws['L4'] = f"OOS (%) [{rng} {bulan[:3]}]"
    ws['N4'] = f"INVENTORY RECAP [{rng} {bulan}]"


def check_data_gaps(list_ds_rows, kd_store_rows, inventory_rows, oos_rows, mat_all_rows, mat_exc_rows, log):
    ld_codes = [r[0] for r in list_ds_rows[1:]]
    sources = {
        'Store Performance (Summary_MTD)': {r[0] for r in kd_store_rows[1:]},
        'Inventory': {r[1] for r in inventory_rows[1:]},
        'OOS': {r[0] for r in oos_rows[1:]},
        'MAT All': {r[1] for r in mat_all_rows[1:]},
        'MAT Exc Beanspot': {r[1] for r in mat_exc_rows[1:]},
    }
    any_gap = False
    for src_name, codes in sources.items():
        missing = [c for c in ld_codes if c not in codes]
        if missing:
            any_gap = True
            log(f"[Peringatan] {len(missing)} toko di List DS tidak ada di {src_name}: {missing} "
                "(akan tampil #N/A di sheet Report)")
    if not any_gap:
        log("Semua toko di List DS ada datanya di 5 sumber lain.")


# ---------------------------------------------------------------- deteksi file
def _match(fn, must_all=(), must_any=(), must_not=()):
    low = fn.lower()
    return (not any(m in low for m in must_not)
            and all(m in low for m in must_all)
            and (not must_any or any(m in low for m in must_any)))


def autodetect_files(paths):
    result = {k: None for k in ['template'] + REQUIRED_KEYS}
    for p in paths:
        if not p.lower().endswith('.xlsx'):
            continue
        fn = os.path.basename(p)
        if _match(fn, ['report'], ['inventory', 'oos', 'mat', 'store performance']):
            result['template'] = p
        elif _match(fn, ['summary_mtd']):
            result['summary_mtd'] = p
        elif _match(fn, ['oos']):
            result['oos'] = p
        elif _match(fn, ['list ds']):
            result['list_ds'] = p
        elif _match(fn, ['inventory_darkstore']):
            result['inventory'] = p
        elif _match(fn, ['mat'], ['excl']):
            result['mat_exc'] = p
        elif _match(fn, ['mat'], must_not=['excl']):
            result['mat_all'] = p
    return result


# ---------------------------------------------------------------- entry point
def run(paths, output_dir, log=None):
    log = log or Log()
    det = autodetect_files(paths)
    missing = [LABELS[k] for k in REQUIRED_KEYS if det[k] is None]
    if missing:
        raise FileNotFoundError("File berikut tidak ditemukan (cek nama file): " + "; ".join(missing))
    if det['template'] is None:
        log("File Report periode sebelumnya tidak ada -> pakai template bawaan.")
        det['template'] = DEFAULT_TEMPLATE
    for k, v in det.items():
        log(f"{k:12s}: {os.path.basename(v)}")

    kd_store_rows = load_kd_store(det['summary_mtd'])
    oos_rows = load_oos(det['oos'])
    list_ds_rows = load_list_ds(det['list_ds'])
    inventory_rows, period_start, period_end = load_inventory(det['inventory'])
    mat_all_rows = load_mat_long(det['mat_all'], 'MAT All')
    mat_exc_rows = load_mat_long(det['mat_exc'], 'MAT Exc Beanspot')
    check_data_gaps(list_ds_rows, kd_store_rows, inventory_rows, oos_rows, mat_all_rows, mat_exc_rows, log)

    wb = openpyxl.load_workbook(det['template'])
    write_table_sheet(wb, 'Store Performance', 'Table8', kd_store_rows)
    write_table_sheet(wb, 'OOS', 'Table6', oos_rows)
    write_table_sheet(wb, 'List DS', 'Table1', list_ds_rows)
    write_table_sheet(wb, 'Inventory', 'Table5', inventory_rows)
    write_table_sheet(wb, 'MAT All', 'Table25', mat_all_rows)
    write_table_sheet(wb, 'MAT Exc Beanspot', 'Table2', mat_exc_rows)
    rebuild_report_sheet(wb, kd_store_rows, list_ds_rows, inventory_rows, period_start, period_end)

    name = (f"REPORT INVENTORY, OOS, MAT, STORE PERFORMANCE "
            f"[MTD {period_end.day} {MONTH_ID[period_end.month - 1][:3]} {period_end.year}].xlsx")
    out = os.path.join(output_dir, name)
    wb.save(out)
    wb.close()
    log(f"Periode {period_start:%d-%m-%Y} s.d. {period_end:%d-%m-%Y}. File dibuat: {name}")
    return out, None, []
