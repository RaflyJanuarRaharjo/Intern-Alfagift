"""Smoke test: bikin data dummy untuk ke-6 laporan lalu jalankan semua modul core.

Jalankan dari root repo:  python tests/smoke_test.py
"""
import datetime as dt
import os
import random
import sys
import tempfile

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from core.common import prepare_inputs  # noqa: E402
from core import (apo_darkstore, darkstore_ds, inventory_report, master_produk, mds_report,  # noqa: E402
                  mtd_performance)

random.seed(1)
STORES = [(f"T{i:03d}", f"DS TOKO {i}", random.choice(["JAKARTA", "BEKASI", "SIDOARJO"])) for i in range(1, 13)]


def make_detail_and_oos(d):
    rows = []
    for day in range(1, 8):
        for kd, nama, cab in STORES:
            if day == 7 and kd == "T003":  # toko tanpa data hari terakhir -> uji fallback
                continue
            sales = random.randint(5_000_000, 50_000_000)
            rows.append(dict(TANGGAL=f"2026-09-{day:02d}", KD_STORE=kd, NAMA_STORE=nama, NAMA_BRANCH=cab,
                             JHK=1, SALES=sales, SALES_TAGI=sales * 0.2, SPD=sales, STD=random.randint(100, 400),
                             APC=random.randint(80_000, 150_000), PERCENT_GM=random.uniform(.1, .2),
                             PCT_OOS_OFMB=random.uniform(0, .1)))
    pd.DataFrame(rows).to_csv(os.path.join(d, "Detail_Data_Sep.csv"), index=False)
    oos = pd.DataFrame([dict(**{"Kode Toko": kd, "Nama Toko": n, "Cabang": c},
                             **{k: random.uniform(0, .1) for k in
                                ["% OOS TAG I", "% OOS OFMB", "% OOS TAG K", "% OOS ALL"]})
                        for kd, n, c in STORES] + [{"Kode Toko": "Grand Total"}])
    oos.to_excel(os.path.join(d, "[OOS] By Toko.xlsx"), index=False)
    rs = pd.DataFrame([dict(KD_STORE=kd, JUMLAH_DELIVERY=100, DELIVERY_ONTIME=random.randint(60, 95),
                            DELIVERY_LATE=5) for kd, _, _ in STORES])
    rs.to_excel(os.path.join(d, "Report Summary Dashboard.xlsx"), index=False)


def _wb_with(path, sheet, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    wb.save(path)


def make_inventory_inputs(d):
    perf_hdr = ['KD_STORE', 'NAMA_STORE', 'KD_BRANCH', 'NAMA_BRANCH', 'REMARK', 'JHK', 'SALES', 'STRUK', 'GM',
                'PERCENT_GM', 'SPD', 'STD', 'APC']
    _wb_with(os.path.join(d, "Summary_MTD_September_2026.xlsx"), "KD_STORE",
             [perf_hdr] + [[kd, n, "B1", c, "DARKSTORE", 7, 1e8, 3000, 1e7, .15, 1.4e7, 400, 1e5]
                           for kd, n, c in STORES])
    _wb_with(os.path.join(d, "[OOS] By Toko.xlsx"), "Sheet 1",
             [["Kode Toko", "Nama Toko", "Cabang", "% OOS TAG I", "% OOS OFMB", "% OOS TAG K", "% OOS ALL"]]
             + [[kd, n, c, .05, .07, .1, .12] for kd, n, c in STORES] + [["Grand Total", None, None, .05]])
    _wb_with(os.path.join(d, "List DS_Covered Invent_Sep.xlsx"), "List DS",
             [["Store Code", "Store Name", "Branch Code", "Branch Name", "Opening Date", "Remark"]]
             + [[kd, n, "B1", c, dt.datetime(2026, 1, 1), "Invent"] for kd, n, c in STORES])
    inv_hdr = ['Branch Code', 'Branch Name', 'Store Code', 'Store Name', 'Item Aktif', 'Item On Stock Total',
               'Item On Stock (OH > 0)', 'Item No Stock (OOS)', 'Item No Stock (OH > 0)', 'Inventory Value',
               'MPKM Value', 'PKM Value', 'N+ Value', 'PKM Exist Value', 'GAP Value (Inv. vs PKM Exist)',
               'DSI PKM', 'DSI OH', 'OOS On Stock', 'OOS % On Stock', 'MAT % On Stock', 'MAT % No Stock',
               'MAT % Beanspot', 'Remark Item']
    inv_rows = [["INVENTORY DARKSTORE"], ["PERIODE 01-09-2026 SD 07-09-2026"], [], inv_hdr]
    for kd, n, c in STORES:
        for remark in ("ALL", "FOOD"):
            inv_rows.append(["B1", c, kd, n, 6000, 5900, 5500, 30, 28, random.randint(10**8, 9 * 10**8), 5e8, 7e8, 0,
                             7e8, 9e7, 16.6, 18.7, 495, .08, .68, .19, .12, remark])
    _wb_with(os.path.join(d, "inventory_darkstore_Sep.xlsx"), "Store", inv_rows)
    for fn, label in (("[MAT] Transaksi All.xlsx", "all"), ("[MAT] Transaksi Excl Beanspot.xlsx", "exc")):
        _wb_with(os.path.join(d, fn), "Sheet 1",
                 [["Kode"] + [s[0] for s in STORES], ["Nama"] + [s[1] for s in STORES],
                  ["01/09"] + [.1] * len(STORES), ["Average"] + [random.uniform(0, .1) for _ in STORES]])


def make_perf_input(d):
    hdr = ["KD_STORE", "NAMA_STORE", "KD_BRANCH", "NAMA_BRANCH", "REMARK", "JHK", "SALES", "STRUK", "GM",
           "PERCENT_GM", "SPD", "STD", "APC", "DELIVERY_ONTIME", "DELIVERY_LATE", "DELIVERY_NO_SLA",
           "JUMLAH_SHIPMENT_TUNDA"]
    _wb_with(os.path.join(d, "perf.xlsx"), "Sheet1",
             [["Performance Darkstore"], [], hdr]
             + [[kd, n, "B1", c, "DS", 7, 1e8, 3000, 1e7, "15.5%", 1.4e7, 400, 1e5,
                 random.randint(50, 95), 5, 2, 1] for kd, n, c in STORES])


def make_master_inputs(d):
    for i in range(2):
        df = pd.DataFrame({"GROUP CODE": [f"G{i}{j}" for j in range(5)], "PLU": range(5),
                           "TAG": ["D", "G", "X", "6", "Z"]})
        with open(os.path.join(d, f"master_{i}.xls"), "w") as f:
            f.write("Report Master Group PLU\n\n")
            df.to_csv(f, sep="\t", index=False)


def make_apo_input(d):
    rows = [dict(store=n, **{"sort status apo": s, "Jumlah Order": random.randint(0, 30)})
            for _, n, _ in STORES for s in apo_darkstore.URUTAN_STATUS]
    pd.DataFrame(rows).to_csv(os.path.join(d, "GLI - NEW SAPA  Performance_ALFAGIFT DARK STORE_Tabel pivot.csv"),
                              index=False)


def make_mds_input(d):
    rows = []
    for day in range(1, 8):
        for kd, nama in mds_report.DEFAULT_MASTER[:8] + [("ZZ99", "BUKAN MDS")]:
            sales = random.randint(1_000_000, 9_000_000)
            rows.append(dict(TANGGAL=f"2026-10-{day:02d}", KD_STORE=kd, NAMA_STORE=nama, KD_BRANCH="B1",
                             NAMA_BRANCH="CABANG", REMARK="DARKSTORE", JHK=1, SALES=sales,
                             SALES_TAGI=sales * .2, SPD=sales, STD=random.randint(50, 200),
                             APC=random.randint(50_000, 90_000), PERCENT_GM=random.uniform(.1, .2)))
    pd.DataFrame(rows).to_csv(os.path.join(d, "Detail_Data_Okt.csv"), index=False)


def make_mds_zip(d):
    """1 ZIP berisi TRX + SLA + Master MDS (+ report lama) seperti pemakaian sebenarnya."""
    import zipfile
    stores = mds_report.DEFAULT_MASTER[:8] + [("ZZ99", "BUKAN MDS")]
    trx, sla = [], []
    for day in range(1, 8):
        for kd, nama in stores:
            sales = random.randint(1_000_000, 9_000_000)
            base = dict(TANGGAL=f"2026-10-{day:02d}", KD_STORE=kd, NAMA_STORE=nama, KD_BRANCH="B1",
                        NAMA_BRANCH="CABANG", REMARK="SAPA_NON_DS", PCT_OOS_OFMB=random.uniform(.01, .05))
            trx.append(dict(base, JHK=1, SALES=sales, SALES_TAGI=sales * .2, SPD=sales,
                            STD=random.randint(50, 200), APC=random.randint(50_000, 90_000),
                            PERCENT_GM=random.uniform(.1, .2)))
            n = random.randint(40, 120)
            late = random.randint(0, 15)
            sla.append(dict(base, JUMLAH_DELIVERY=n, DELIVERY_ONTIME=n - late, DELIVERY_LATE=late))
    pd.DataFrame(trx).to_csv(os.path.join(d, "Detail_Data_Okt TRX.csv"), index=False)
    pd.DataFrame(sla).to_csv(os.path.join(d, "Detail_Data_Okt SLA.csv"), index=False)
    pd.DataFrame(mds_report.DEFAULT_MASTER[:8], columns=["KD STORE", "NAMA STORE"]).assign(
        **{"REMARK MDS": "MINI DARKSTORE"}).to_excel(os.path.join(d, "Master MDS.xlsx"), index=False)
    pd.DataFrame({"x": [1]}).to_excel(os.path.join(d, "Report Daily MDS - lama.xlsx"), index=False)
    zp = os.path.join(d, "Data Report Daily MDS.zip")
    with zipfile.ZipFile(zp, "w") as z:
        for f in ("Detail_Data_Okt TRX.csv", "Detail_Data_Okt SLA.csv", "Master MDS.xlsx",
                  "Report Daily MDS - lama.xlsx"):
            z.write(os.path.join(d, f), "Data Report Daily MDS/" + f)
            os.remove(os.path.join(d, f))


def ls(d):
    return [os.path.join(d, f) for f in sorted(os.listdir(d))]


def main():
    root = tempfile.mkdtemp()
    out = os.path.join(root, "out")
    os.makedirs(out)
    dirs = {k: os.path.join(root, k) for k in ("ds", "inv", "perf", "mp", "apo", "mds")}
    for d in dirs.values():
        os.makedirs(d)
    make_detail_and_oos(dirs["ds"])
    make_inventory_inputs(dirs["inv"])
    make_perf_input(dirs["perf"])
    make_master_inputs(dirs["mp"])
    make_apo_input(dirs["apo"])
    make_mds_input(dirs["mds"])
    os.makedirs(os.path.join(root, "mdszip"))
    make_mds_zip(os.path.join(root, "mdszip"))

    ds_files = [p for p in ls(dirs["ds"]) if "Report Summary" not in p]
    results = {
        "daily_ds": darkstore_ds.run(ds_files, out),
        "mtd_ds": darkstore_ds.run(ls(dirs["ds"]), out, with_delivery=True),
        "inventory": inventory_report.run(ls(dirs["inv"]), out),
        "mtd_perf": mtd_performance.run(ls(dirs["perf"]), out, tanggal_laporan="7 SEP 2026"),
        "master_produk": master_produk.run(ls(dirs["mp"]), out),
        "mds": mds_report.run(ls(dirs["mds"]), out),
        "mds_zip": mds_report.run(prepare_inputs(os.path.join(root, "mdszip")),
                                  os.path.join(root, "out_zip") if os.makedirs(os.path.join(root, "out_zip"), exist_ok=True) is None else out),
        "apo": apo_darkstore.run(ls(dirs["apo"]), out, tanggal="07/10/2026", jam="10.15"),
    }
    for name, (path, _preview, extras) in results.items():
        assert os.path.getsize(path) > 0, name
        wb = openpyxl.load_workbook(path)
        print(f"OK  {name:14s} -> {os.path.basename(path)}  sheets={wb.sheetnames}  extras={len(extras)}")
    print("Output:", out)
    return out


if __name__ == "__main__":
    main()
